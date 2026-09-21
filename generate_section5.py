import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

md_path = r"C:\Users\Usuario\.gemini\antigravity\brain\a5af4022-b119-45ee-bcea-888007b1c989\.user_uploaded\media_1789924561817.md"

with open(md_path, 'r', encoding='utf-8') as f:
    text = f.read()

# 1. Parse the Catalog (5.1)
catalog = {}
catalog_match = re.search(r'## \*\*5\.1 Catálogo general\*\*.*?(?=\n## \*\*5\.2)', text, re.DOTALL | re.IGNORECASE)
if not catalog_match:
    catalog_match = re.search(r'5\.1 Cat.*?(?=\n5\.2)', text, re.DOTALL | re.IGNORECASE)

if catalog_match:
    for line in catalog_match.group(0).split('\n'):
        if '|' in line and 'RF-' in line:
            parts = [p.strip() for p in line.split('|')]
            if len(parts) >= 8:
                rf_id = parts[1]
                if rf_id.startswith('RF-'):
                    catalog[rf_id] = {
                        'nombre': parts[2],
                        'modulo': parts[3],
                        'actor': parts[4],
                        'prioridad': parts[5],
                        'estado': parts[6],
                        'version': parts[7]
                    }

# Update Catalog with Stripe, Workflows, Firmas
catalog_updates = {
    'RF-062': {'nombre': 'Integrar Stripe SDK modo Test (suscripción, upgrade, downgrade, cancelación)', 'modulo': 'M-11 Planes y Facturación'},
    'RF-063': {'nombre': 'Simular confirmación de pago en Stripe y activar plan', 'modulo': 'M-11 Planes y Facturación'},
    'RF-067': {'nombre': 'Landing pública en / para visitantes sin sesión', 'modulo': 'M-14 Sitio Público'},
    'RF-068': {'nombre': 'Redirección desde landing según sesión', 'modulo': 'M-14 Sitio Público'},
    'RF-069': {'nombre': 'Registro extendido: paso datos organización', 'modulo': 'M-15 Registro con Plan'},
    'RF-070': {'nombre': 'Paso selección de plan con comparativa', 'modulo': 'M-15 Registro con Plan'},
    'RF-071': {'nombre': 'Plan de pago -> checkout Stripe directo; Demo -> Demo', 'modulo': 'M-15 Registro con Plan'},
    'RF-072': {'nombre': 'Crear Workflow en documento/lotes (draft, in_progress)', 'modulo': 'M-12 Workflows', 'actor': 'Admin/Miembro', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'},
    'RF-073': {'nombre': 'Asignar roles/usuarios a nodos del flujo (linear, parallel)', 'modulo': 'M-12 Workflows', 'actor': 'Admin/Miembro', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'},
    'RF-074': {'nombre': 'Posicionamiento de firmas visuales en coordenadas X,Y', 'modulo': 'M-12 Workflows', 'actor': 'Admin/Miembro', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'},
    'RF-075': {'nombre': 'Restringir descargas parciales/completas en workflow', 'modulo': 'M-12 Workflows', 'actor': 'Admin/Miembro', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'},
    'RF-076': {'nombre': 'Aprobar, rechazar o firmar documento según nodo asignado', 'modulo': 'M-12 Workflows', 'actor': 'Revisor/Firmante', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'},
    'RF-077': {'nombre': 'Estampar firma visual en PDF usando canvas y pdf-lib', 'modulo': 'M-13 Firmas PKI', 'actor': 'Firmante', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'},
    'RF-078': {'nombre': 'Firma criptográfica con node-forge y hash inmutable', 'modulo': 'M-13 Firmas PKI', 'actor': 'Sistema', 'prioridad': 'Alta', 'estado': 'Aprobado', 'version': 'MVP'}
}

for k, v in catalog_updates.items():
    if k in catalog:
        catalog[k].update(v)
    else:
        catalog[k] = v

# 2. Parse Fichas (5.2)
fichas = {}
fichas_match = re.search(r'## \*\*5\.2 Ficha detallada.*?(?=## \*\*5\.3|# \*\*6\.)', text, re.DOTALL | re.IGNORECASE)
if not fichas_match:
    fichas_match = re.search(r'5\.2 Ficha detallada.*?(?=5\.3|6\.)', text, re.DOTALL | re.IGNORECASE)

field_names = [
    'ID y nombre', 'Descripción', 'Fuente', 'Justificación', 'Actor principal',
    'Precondiciones', 'Disparador', 'Entradas', 'Procesamiento / comportamiento',
    'Salidas', 'Postcondiciones', 'Excepciones', 'Reglas de negocio asociadas',
    'Dependencias', 'Prioridad', 'Versión objetivo', 'Criterios de aceptación',
    'Método de verificación', 'Responsable de validación', 'Estado'
]

if fichas_match:
    blocks = fichas_match.group(0).split('Campo | Contenido')
    for block in blocks[1:]:
        lines = [l for l in block.split('\n') if '|' in l and '---' not in l]
        current_ficha = {f: "No especificado" for f in field_names}
        rf_id = None
        for line in lines:
            parts = [p.strip() for p in line.split('|')]
            if len(parts) >= 3:
                campo = parts[1].replace('*', '').strip()
                contenido = parts[2].replace('\\-', '-').strip()
                
                norm_campo = campo.lower().replace('ó', 'o').replace('í', 'i')
                matched_field = None
                for fn in field_names:
                    if fn.lower().replace('ó', 'o').replace('í', 'i') in norm_campo:
                        matched_field = fn
                        break
                
                if matched_field:
                    current_ficha[matched_field] = contenido
                    if matched_field == 'ID y nombre' and 'RF-' in contenido:
                        m = re.search(r'RF-\d{3}', contenido)
                        if m:
                            rf_id = m.group(0)
        if rf_id:
            if rf_id not in fichas:
                fichas[rf_id] = current_ficha
            else:
                if 'Stripe' in current_ficha.get('Descripción', '') and 'PayPal' in fichas[rf_id].get('Descripción', ''):
                    fichas[rf_id] = current_ficha

new_fichas = {
    'RF-072': {
        'ID y nombre': 'RF-072 - Crear Workflow en documento/lotes (draft, in_progress)',
        'Descripción': 'El sistema debe permitir asociar un flujo de revisión/firma a un documento o lote de documentos.',
        'Fuente': 'STK-05 / OBJ-05 / HU-33 / CU-33',
        'Justificación': 'Automatización de revisión documental (M-12).',
        'Actor principal': 'Admin/Miembro',
        'Precondiciones': '- Usuario tiene permisos en la sala. - Documento existe.',
        'Disparador': 'Usuario crea flujo en interfaz de documento.',
        'Entradas': 'Selección de documento y nodos de flujo.',
        'Procesamiento / comportamiento': 'Insertar registro en document_workflows (estado draft).',
        'Salidas': 'Workflow en BD.',
        'Postcondiciones': 'Flujo listo para firmantes.',
        'Excepciones': '1a. Documento ya tiene flujo -> error.',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-07 Recursos',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'Se crea y asocia a la org.',
        'Método de verificación': 'Prueba funcional',
        'Responsable de validación': 'QA',
        'Estado': 'Aprobado'
    },
    'RF-073': {
        'ID y nombre': 'RF-073 - Asignar roles/usuarios a nodos del flujo (linear, parallel)',
        'Descripción': 'Asignar Revisor o Firmante a los nodos del workflow.',
        'Fuente': 'STK-05 / OBJ-05 / HU-34 / CU-34',
        'Justificación': 'Determina quién debe aprobar/firmar el documento.',
        'Actor principal': 'Admin/Miembro',
        'Precondiciones': '- Workflow en estado draft.',
        'Disparador': 'Usuario añade un nodo al flujo.',
        'Entradas': 'Usuario destino, rol del nodo.',
        'Procesamiento / comportamiento': 'Insertar en workflow_nodes.',
        'Salidas': 'Nodos asignados.',
        'Postcondiciones': 'El documento requiere la aprobación del usuario.',
        'Excepciones': '1a. Usuario no existe -> error.',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-12 Workflows',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'Usuarios se agregan correctamente a nodos.',
        'Método de verificación': 'Prueba funcional',
        'Responsable de validación': 'QA',
        'Estado': 'Aprobado'
    },
    'RF-074': {
        'ID y nombre': 'RF-074 - Posicionamiento de firmas visuales en coordenadas X,Y',
        'Descripción': 'Permitir definir las posiciones exactas (página, X, Y) donde se estampará la firma visual.',
        'Fuente': 'STK-05 / OBJ-05 / HU-35 / CU-35',
        'Justificación': 'Cumplimiento de representación visual de firma.',
        'Actor principal': 'Admin/Miembro',
        'Precondiciones': '- Workflow con al menos un nodo firmante.',
        'Disparador': 'Usuario abre previsualización de documento para colocar firmas.',
        'Entradas': 'Coordenadas (X,Y) y página.',
        'Procesamiento / comportamiento': 'Guardar en workflow_signature_positions.',
        'Salidas': 'Posiciones registradas en BD.',
        'Postcondiciones': 'Al firmar, la firma aparecerá en esa posición.',
        'Excepciones': '1a. Coordenadas fuera de límite -> ajustar al borde.',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-12 Workflows',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'Se registran posiciones visuales por usuario.',
        'Método de verificación': 'Prueba funcional',
        'Responsable de validación': 'QA',
        'Estado': 'Aprobado'
    },
    'RF-075': {
        'ID y nombre': 'RF-075 - Restringir descargas parciales/completas en workflow',
        'Descripción': 'Bloquear la descarga del documento hasta que se complete el flujo (configurable).',
        'Fuente': 'STK-05 / OBJ-05 / HU-36 / CU-36',
        'Justificación': 'Evitar filtraciones de documentos no aprobados.',
        'Actor principal': 'Admin/Miembro',
        'Precondiciones': '- Flujo in_progress.',
        'Disparador': 'Usuario intenta descargar.',
        'Entradas': 'Petición de descarga.',
        'Procesamiento / comportamiento': 'Validar estado del flujo y permisos. Rechazar si está restringido.',
        'Salidas': 'Archivo o HTTP 403.',
        'Postcondiciones': '-',
        'Excepciones': '1a. Admin solicita descarga -> permitir.',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-12 Workflows',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'Usuarios sin permiso no pueden descargar el documento.',
        'Método de verificación': 'Prueba funcional',
        'Responsable de validación': 'QA',
        'Estado': 'Aprobado'
    },
    'RF-076': {
        'ID y nombre': 'RF-076 - Aprobar, rechazar o firmar documento según nodo asignado',
        'Descripción': 'Permitir al usuario asignado aprobar, rechazar o firmar el documento.',
        'Fuente': 'STK-05 / OBJ-05 / HU-37 / CU-37',
        'Justificación': 'Acción principal del ciclo de vida del documento.',
        'Actor principal': 'Revisor/Firmante',
        'Precondiciones': '- Es el turno del usuario en el flujo.',
        'Disparador': 'Usuario pulsa Aprobar, Rechazar o Firmar.',
        'Entradas': 'Confirmación de acción.',
        'Procesamiento / comportamiento': 'Actualizar estado del nodo. Si es firma, ejecutar RF-077/078.',
        'Salidas': 'Nodo completado.',
        'Postcondiciones': 'El flujo avanza al siguiente nodo o finaliza.',
        'Excepciones': '1a. Rechazo -> flujo finaliza en cancelled.',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-12 Workflows',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'El estado del nodo se actualiza correctamente y avanza el flujo.',
        'Método de verificación': 'Prueba funcional',
        'Responsable de validación': 'QA',
        'Estado': 'Aprobado'
    },
    'RF-077': {
        'ID y nombre': 'RF-077 - Estampar firma visual en PDF usando canvas y pdf-lib',
        'Descripción': 'El sistema debe dibujar la firma del usuario sobre el PDF usando las coordenadas guardadas.',
        'Fuente': 'STK-05 / OBJ-05 / HU-38 / CU-38',
        'Justificación': 'Representación visual de la firma legal.',
        'Actor principal': 'Firmante',
        'Precondiciones': '- Nodo de firma activado.',
        'Disparador': 'Usuario firma el documento.',
        'Entradas': 'Firma generada en canvas.',
        'Procesamiento / comportamiento': 'pdf-lib inserta la imagen en las coordenadas (X,Y).',
        'Salidas': 'PDF modificado visualmente.',
        'Postcondiciones': 'El documento contiene la estampa.',
        'Excepciones': '-',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-13 Firmas PKI',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'PDF actualizado con la imagen de la firma.',
        'Método de verificación': 'Prueba funcional',
        'Responsable de validación': 'QA',
        'Estado': 'Aprobado'
    },
    'RF-078': {
        'ID y nombre': 'RF-078 - Firma criptográfica con node-forge y hash inmutable',
        'Descripción': 'Calcular hash del PDF modificado y aplicar firma criptográfica PKI.',
        'Fuente': 'STK-05 / OBJ-05 / HU-39 / CU-39',
        'Justificación': 'Inmutabilidad y validación legal.',
        'Actor principal': 'Sistema',
        'Precondiciones': '- Firma visual aplicada.',
        'Disparador': 'Finalización del estampado visual.',
        'Entradas': 'PDF buffer.',
        'Procesamiento / comportamiento': 'node-forge calcula SHA256 y firma con clave privada. Guarda en document_signatures.',
        'Salidas': 'Registro inmutable en BD.',
        'Postcondiciones': 'El documento queda sellado.',
        'Excepciones': '-',
        'Reglas de negocio asociadas': 'RN-010',
        'Dependencias': 'M-13 Firmas PKI',
        'Prioridad': 'Must (Alta)',
        'Versión objetivo': 'MVP',
        'Criterios de aceptación': 'Cualquier modificación posterior al PDF invalida el hash.',
        'Método de verificación': 'Prueba técnica y forense',
        'Responsable de validación': 'QA / Dev',
        'Estado': 'Aprobado'
    }
}

for k, v in new_fichas.items():
    fichas[k] = v

for k in ['RF-062', 'RF-063', 'RF-071']:
    if k in fichas:
        fichas[k]['ID y nombre'] = fichas[k]['ID y nombre'].replace('PayPal', 'Stripe')
        fichas[k]['Descripción'] = fichas[k]['Descripción'].replace('PayPal SDK Sandbox', 'Stripe SDK modo Test')
        fichas[k]['Descripción'] = fichas[k]['Descripción'].replace('PayPal', 'Stripe')
        fichas[k]['Procesamiento / comportamiento'] = fichas[k]['Procesamiento / comportamiento'].replace('PayPal', 'Stripe')
        fichas[k]['Criterios de aceptación'] = fichas[k]['Criterios de aceptación'].replace('PayPal', 'Stripe')

sorted_keys = [f"RF-{str(i).zfill(3)}" for i in range(1, 79)]

doc = Document()
doc.add_heading('Sección 5 Corregida', 0)

doc.add_heading('5. Requisitos funcionales', level=1)
doc.add_heading('5.1 Catálogo general', level=2)

table1 = doc.add_table(rows=1, cols=7)
table1.style = 'Table Grid'
hdr_cells = table1.rows[0].cells
hdr_cells[0].text = 'ID'
hdr_cells[1].text = 'Nombre / Descripción corta'
hdr_cells[2].text = 'Módulo'
hdr_cells[3].text = 'Actor principal'
hdr_cells[4].text = 'Prioridad'
hdr_cells[5].text = 'Estado'
hdr_cells[6].text = 'Versión'

for k in sorted_keys:
    if k in catalog:
        row = table1.add_row().cells
        row[0].text = k
        row[1].text = catalog[k]['nombre']
        row[2].text = catalog[k]['modulo']
        row[3].text = catalog[k]['actor']
        row[4].text = catalog[k]['prioridad']
        row[5].text = catalog[k]['estado']
        row[6].text = catalog[k]['version']

doc.add_page_break()
doc.add_heading('5.2 Ficha detallada de requisito funcional', level=2)

for k in sorted_keys:
    if k in fichas:
        doc.add_heading(f'Ficha: {k}', level=3)
        t = doc.add_table(rows=1, cols=2)
        t.style = 'Table Grid'
        t.rows[0].cells[0].text = 'Campo'
        t.rows[0].cells[1].text = 'Contenido'
        
        for field in field_names:
            r = t.add_row()
            r.cells[0].text = field
            r.cells[1].text = fichas[k].get(field, 'No especificado')
        
        doc.add_paragraph()

doc.save(r"c:\Users\Usuario\Desktop\backroom proyecto y documentacion\cambios-documentos\Seccion_5_Corregida.docx")
print("DOCX generado correctamente en cambios-documentos\Seccion_5_Corregida.docx")
