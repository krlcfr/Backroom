from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()
doc.add_heading('Guía de Cambios: Realidad del Proyecto vs Documento', 0)

doc.add_paragraph(
    "Tras analizar exhaustivamente el código fuente del proyecto (app/api/workflows, package.json, y componentes), esta es la realidad técnica:\n"
    "1. Sí existen MP3 y MP4: El proyecto usa <video> y <audio> de HTML5 y los campos `recursos_tipo_check` en la BD permiten mp3 y mp4. NO debes quitarlos.\n"
    "2. Sí existen Workflows y Firmas PKI: El código usa `node-forge` para firmas criptográficas inmutables y `pdf-lib` para estampas visuales. Las rutas de workflows son completamente reales. NO debes quitar los RF-072 al RF-078.\n"
    "3. No existe PayPal: La integración real y oficial en el código es con Stripe."
)

doc.add_heading('1. Cambios a realizar en 5.1 (Catálogo general)', level=1)
doc.add_paragraph('Ubicación: Al final de la tabla 5.1 que me pasaste.')
doc.add_paragraph('Acción: AÑADIR las siguientes filas (del 072 al 078), ya que se te olvidó agregarlas en tu Word y SÍ son reales.')

table = doc.add_table(rows=1, cols=7)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = 'ID'
hdr[1].text = 'Nombre / Descripción corta'
hdr[2].text = 'Módulo'
hdr[3].text = 'Actor'
hdr[4].text = 'Prioridad'
hdr[5].text = 'Estado'
hdr[6].text = 'Versión'

nuevos_rf = [
    ('RF-072', 'Crear Workflow en documento/lotes (draft, in_progress)', 'M-12 Workflows', 'Admin/Miembro', 'Alta', 'Aprobado', 'MVP'),
    ('RF-073', 'Asignar roles/usuarios a nodos del flujo (linear, parallel)', 'M-12 Workflows', 'Admin/Miembro', 'Alta', 'Aprobado', 'MVP'),
    ('RF-074', 'Posicionamiento de firmas visuales en coordenadas X,Y', 'M-12 Workflows', 'Admin/Miembro', 'Alta', 'Aprobado', 'MVP'),
    ('RF-075', 'Restringir descargas parciales/completas en workflow', 'M-12 Workflows', 'Admin/Miembro', 'Alta', 'Aprobado', 'MVP'),
    ('RF-076', 'Aprobar, rechazar o firmar documento según nodo asignado', 'M-12 Workflows', 'Revisor', 'Alta', 'Aprobado', 'MVP'),
    ('RF-077', 'Estampar firma visual en PDF usando canvas y pdf-lib', 'M-13 Firmas PKI', 'Firmante', 'Alta', 'Aprobado', 'MVP'),
    ('RF-078', 'Firma criptográfica con node-forge y hash inmutable', 'M-13 Firmas PKI', 'Sistema', 'Alta', 'Aprobado', 'MVP')
]
for rf in nuevos_rf:
    row = table.add_row().cells
    for i in range(7):
        row[i].text = rf[i]

doc.add_heading('2. Cambios a realizar en 5.2 (Fichas detalladas)', level=1)

doc.add_heading('2.1 Modificar Fichas de Pagos (Quitar PayPal)', level=2)
doc.add_paragraph('Ubicación: En las fichas RF-062, RF-063 y RF-071.')
doc.add_paragraph('Acción: MODIFICAR y cambiar cualquier mención de "PayPal" por "Stripe". Te dejo el texto exacto a cambiar:')

table2 = doc.add_table(rows=1, cols=3)
table2.style = 'Table Grid'
hdr2 = table2.rows[0].cells
hdr2[0].text = 'Ficha'
hdr2[1].text = 'Campo a cambiar'
hdr2[2].text = 'Nuevo contenido (realidad)'

cambios = [
    ('RF-062', 'ID y nombre', 'RF-062 - Integrar Stripe SDK modo Test (suscripción, upgrade, downgrade, cancelación)'),
    ('RF-062', 'Descripción', 'El sistema debe integrar Stripe SDK en modo Test para el flujo de suscripción.'),
    ('RF-062', 'Procesamiento', 'Redirigir a Stripe Checkout Session (Test). Guardar stripe_subscription_id.'),
    ('RF-062', 'Criterios', 'Integración y redirección correcta a Stripe Sandbox.'),
    ('RF-063', 'ID y nombre', 'RF-063 - Simular confirmación de pago en Stripe y activar plan'),
    ('RF-063', 'Criterios', 'Simulación de Stripe Sandbox exitosa. Límites aplicados.'),
    ('RF-071', 'ID y nombre', 'RF-071 - Plan de pago -> checkout Stripe directo')
]
for cb in cambios:
    r = table2.add_row().cells
    r[0].text = cb[0]
    r[1].text = cb[1]
    r[2].text = cb[2]

doc.add_heading('2.2 Añadir fichas faltantes de Workflows', level=2)
doc.add_paragraph('Ubicación: Al final del documento (después del RF-071).')
doc.add_paragraph('Acción: AÑADIR las 7 fichas completas del RF-072 al RF-078.')
doc.add_paragraph('NOTA: Como necesitas las tablas con las 20 filas (Precondiciones, Entradas, etc.), por favor cópialas y pégalas directamente desde el archivo "Seccion_5_Corregida.docx" que generé en el paso anterior. ¡Ahí están perfectas y listas para el profesor!')

doc.save(r"c:\Users\Usuario\Desktop\backroom proyecto y documentacion\cambios-documentos\Cambios_Pendientes_Realidad_Proyecto.docx")
